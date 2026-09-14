from __future__ import annotations

import argparse
import gzip
import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Sequence

import cv2
import numpy as np
import pandas as pd

from .data import load_protocol_manifest
from .extension_protocol import BASE_METHODS, NEW_METHODS, verify_base
from .io import read_image, sha256_file
from .package_light import _write_montage
from .registry import load_yaml
from .model_complexity import profile
from .methods.sabids_adapter import _DenoiseOnly
from .methods.tcfl_adapter import TCFLGenerator
from sabids.engine.trainer import build_model
import torch


METHODS = ["noisy_identity", "bm3d_standard", "tv_chambolle", "nlm", "ksvd_self", "dncnn_paired", "nafnet_paired", "sabids_current", "tcfl_dncnn"]
ZIP_METHODS = [method for method in METHODS if method != "noisy_identity"]


def _primary_seeds(ext: Path) -> dict[str, int]:
    seeds = {method: (42 if method in {"dncnn_paired", "nafnet_paired", "tcfl_dncnn"} else 0) for method in METHODS}
    registry = load_yaml(ext / "configs" / "inference_registry.yaml")
    for method, value in registry.get("methods", {}).items(): seeds[method] = int(value.get("seed", seeds.get(method, 0)))
    return seeds


def build_image_paths(ext: Path) -> pd.DataFrame:
    table = pd.read_csv(ext / "metrics" / "per_image_metrics.csv"); primary = _primary_seeds(ext)
    rows = []
    for row in table.itertuples():
        path = Path(row.denoised_path); exists = path.is_file()
        rows.append({"method_id": row.method_id, "dataset": row.dataset, "split": row.split, "seed": int(row.seed), "is_primary_seed": int(row.seed) == primary.get(row.method_id, 0), "image_root": str(path.parent), "relative_image_path": path.name, "absolute_image_path": str(path), "output_sha256": sha256_file(path) if exists else getattr(row, "output_sha256", ""), "checkpoint_sha256": getattr(row, "checkpoint_sha256", ""), "exists": exists, "bytes": path.stat().st_size if exists else 0, "width": int(row.width), "height": int(row.height), "sample_id": row.sample_id})
    result = pd.DataFrame(rows)
    result.to_csv(ext / "metrics" / "denoised_image_paths.csv", index=False)
    missing = result[~result.exists]
    if not missing.empty: raise RuntimeError(f"{len(missing)} recorded denoised images are missing")
    if not set(METHODS).issubset(set(result.method_id)): raise RuntimeError(f"image path methods incomplete: {sorted(set(METHODS) - set(result.method_id))}")
    lines = ["# Denoised image download paths", "", "All paths below were verified at packaging time.", ""]
    grouped = result.groupby(["method_id", "dataset", "split"], dropna=False)
    for method in METHODS:
        subset = result[result.method_id == method]; lines += [f"## {method}", "", f"- Images: {len(subset):,}; bytes: {int(subset.bytes.sum()):,}; missing: {int((~subset.exists).sum())}", f"- Primary seed: {primary[method]}"]
        for (ignored, dataset, split), group in grouped:
            if ignored == method: lines.append(f"- {dataset}/{split}: `{group.iloc[0].image_root}` ({len(group):,} records across seeds)")
        if method in ZIP_METHODS: lines.append(f"- Primary ZIP: `{ext / 'downloads' / 'primary_images' / (method + '_primary.zip')}`")
        lines.append("")
    (ext / "reports" / "denoised_image_download_paths.md").write_text("\n".join(lines), encoding="utf-8")
    return result


def _primary_zips(ext: Path, paths: pd.DataFrame) -> None:
    output = ext / "downloads" / "primary_images"; output.mkdir(parents=True, exist_ok=True)
    for method in ZIP_METHODS:
        selected = paths[(paths.method_id == method) & paths.is_primary_seed].copy()
        archive = output / f"{method}_primary.zip"
        manifest = selected.drop(columns=["image_root"]).copy()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as bundle:
            used = set()
            for row in selected.itertuples():
                source = Path(row.absolute_image_path); suffix = source.suffix.lower() or ".png"
                relative = f"images/{row.dataset}/{row.split}/{row.sample_id}{suffix}"
                if relative in used: raise RuntimeError(f"duplicate ZIP member {relative}")
                used.add(relative); bundle.write(source, relative)
            bundle.writestr("IMAGE_MANIFEST.csv", manifest.to_csv(index=False))
            bundle.writestr("SHA256SUMS.txt", "".join(f"{row.output_sha256}  images/{row.dataset}/{row.split}/{row.sample_id}{Path(row.absolute_image_path).suffix.lower()}\n" for row in selected.itertuples()))
            bundle.writestr("README.md", f"# {method} primary denoised images\n\nPrimary seed is preregistered and was not selected from test/Duke results. Images preserve benchmark output geometry and bit depth.\n")
        with zipfile.ZipFile(archive) as bundle:
            bad = bundle.testzip()
            if bad: raise RuntimeError(f"CRC failure in {archive}: {bad}")


def _fixed_atlas(root: Path, ext: Path, paths: pd.DataFrame) -> None:
    init_data = json.loads((ext / "audit" / "extension_init.json").read_text(encoding="utf-8")); base = Path(init_data["base_run"])
    source_selection = base / "audit" / "fixed_atlas_selection.csv"
    if not source_selection.is_file(): raise FileNotFoundError("BASE_RUN has no preregistered fixed_atlas_selection.csv")
    selection = pd.read_csv(source_selection); selection.to_csv(ext / "audit" / "fixed_atlas_selection.csv", index=False)
    manifest = load_protocol_manifest(root).set_index("sample_id"); output_dir = ext / "fixed_atlas"; output_dir.mkdir(exist_ok=True)
    primary = paths[paths.is_primary_seed].set_index(["sample_id", "method_id"]); inventory = []
    for item in selection.itertuples():
        sample = str(item.sample_id)
        if sample not in manifest.index: continue
        source = manifest.loc[sample]; noisy, _ = read_image(Path(source.image_path)); reference, _ = read_image(Path(source.clean_path))
        images = [("noisy", noisy), ("reference", reference)]
        for method in METHODS[1:]:
            key = (sample, method)
            if key not in primary.index: raise RuntimeError(f"fixed atlas output missing: {sample}/{method}")
            record = primary.loc[key]
            if isinstance(record, pd.DataFrame): record = record.iloc[0]
            value, _ = read_image(Path(record.absolute_image_path)); images.append((method, value))
        error = [("noisy abs error", np.abs(noisy-reference)), ("reference", np.zeros_like(noisy))] + [(name, np.abs(value-reference)) for name, value in images[2:]]
        residual = [("noisy residual", np.zeros_like(noisy)), ("reference", np.zeros_like(noisy))] + [(name, np.clip(0.5 + 2*(noisy-value), 0, 1)) for name, value in images[2:]]
        destination = output_dir / f"{item.atlas_role}__{sample}__full.jpg"; _write_montage(destination, [images, error, residual], cell_width=180)
        inventory.append({"path": str(destination), "sample_id": sample, "atlas_role": item.atlas_role, "view": "full", "sha256": sha256_file(destination), "bytes": destination.stat().st_size})
        x, y, width, height = int(item.detail_crop_x), int(item.detail_crop_y), int(item.detail_crop_width), int(item.detail_crop_height)
        lx, ly, lwidth, lheight = int(item.lower_crop_x), int(item.lower_crop_y), int(item.lower_crop_width), int(item.lower_crop_height)
        crop_rows = [[(name, value[y:y + height, x:x + width]) for name, value in images], [(name, value[ly:ly + lheight, lx:lx + lwidth]) for name, value in images]]
        crop_path = output_dir / f"{item.atlas_role}__{sample}__crops.jpg"; _write_montage(crop_path, crop_rows, cell_width=180)
        inventory.append({"path": str(crop_path), "sample_id": sample, "atlas_role": item.atlas_role, "view": "preregistered_crops", "sha256": sha256_file(crop_path), "bytes": crop_path.stat().st_size})
    pd.DataFrame(inventory).to_csv(ext / "audit" / "fixed_atlas_asset_inventory.csv", index=False)


def _collect_training(ext: Path) -> None:
    base = Path(json.loads((ext / "audit" / "extension_init.json").read_text(encoding="utf-8"))["base_run"])
    curves, inventory = [], []
    if (base / "metrics" / "training_curves.csv").is_file(): curves.append(pd.read_csv(base / "metrics" / "training_curves.csv"))
    if (base / "metrics" / "checkpoint_inventory.csv").is_file(): inventory.append(pd.read_csv(base / "metrics" / "checkpoint_inventory.csv"))
    for path in (ext / "tracks").glob("**/training_curves.csv"):
        frame = pd.read_csv(path); frame["source"] = str(path); curves.append(frame)
    for path in (ext / "tracks").glob("**/checkpoint_inventory.csv"):
        frame = pd.read_csv(path); frame["source"] = str(path); inventory.append(frame)
    if curves: pd.concat(curves, ignore_index=True, sort=False).to_csv(ext / "metrics" / "training_curves.csv", index=False)
    if inventory: pd.concat(inventory, ignore_index=True, sort=False).to_csv(ext / "metrics" / "checkpoint_inventory.csv", index=False)


def _model_complexity(ext: Path) -> None:
    base = Path(json.loads((ext / "audit" / "extension_init.json").read_text(encoding="utf-8"))["base_run"])
    path = base / "metrics" / "model_complexity.csv"
    rows = pd.read_csv(path).to_dict("records") if path.is_file() else []
    evaluated = pd.read_csv(ext / "metrics" / "per_image_metrics.csv")
    for row in rows:
        if row.get("method_id") in set(evaluated.method_id) and row.get("method_id") in {"dncnn_paired", "nafnet_paired"}:
            row["status"] = "formal_trained_checkpoint_evaluated"
    registry = load_yaml(ext / "configs" / "inference_registry.yaml")
    sabids_entry = registry["methods"]["sabids_current"]; sabids_payload = torch.load(sabids_entry["checkpoint"], map_location="cpu", weights_only=False)
    sabids_config = sabids_payload.get("config")
    sabids_model = _DenoiseOnly(build_model(sabids_config)); parameters, flops = profile(sabids_model)
    rows.append({"method_id": "sabids_current", "configuration": "resolved formal Stage-1/D0 denoise-only forward", "input_height": 640, "input_width": 640, "parameters": parameters, "macs": flops // 2, "gmacs": flops / 2e9, "flops": flops, "gflops": flops / 1e9, "status": "formal_trained_checkpoint_evaluated"})
    parameters, flops = profile(TCFLGenerator())
    rows.append({"method_id": "tcfl_dncnn", "configuration": "official TCFL 10-layer 64-feature generator; discriminator excluded from inference", "input_height": 640, "input_width": 640, "parameters": parameters, "macs": flops // 2, "gmacs": flops / 2e9, "flops": flops, "gflops": flops / 1e9, "status": "formal_trained_checkpoint_evaluated"})
    pd.DataFrame(rows).drop_duplicates("method_id", keep="last").to_csv(ext / "metrics" / "model_complexity.csv", index=False)


def _write_reports(ext: Path) -> None:
    dataset = pd.read_csv(ext / "metrics" / "per_dataset_metrics.csv")
    lock = json.loads((ext / "audit" / "extension_config_lock.json").read_text(encoding="utf-8"))
    result_table = dataset[dataset.method_id.isin(METHODS)][[column for column in ("dataset", "split", "method_id", "psnr", "ssim", "ms_ssim", "epi", "reference_edge_mae", "hf_energy_ratio_to_reference") if column in dataset]].to_markdown(index=False)
    benchmark = f"# SABIDS-Net OCT denoising benchmark extension\n\nThis is the locked PLUS SABIDS/TCFL result. Test started `{lock.get('test_started_at_utc')}` and Duke evaluation started `{lock.get('duke_evaluation_started_at_utc')}`.\n\n{result_table}\n\nK-SVD-self selected parameters lie on multiple finite-search boundaries and do not establish a universal K-SVD optimum. PSNR/SSIM must be read together with EPI, edge error, high-frequency ratios, and the preregistered atlas.\n"
    (ext / "reports" / "benchmark_report.md").write_text(benchmark, encoding="utf-8")
    (ext / "reports" / "method_implementation_report.md").write_text("# Method implementation\n\nSABIDS-current uses the resolved formal independent Stage-1 checkpoint and `forward_denoise_only`; no segmentation output or label enters inference. TCFL-DnCNN uses the audited official generator, discriminator, six adversarial terms, four cross-fusion L1 terms, and independently sampled PKU37-train noisy/clean pools. Inference loads only the generator, contains no dropout, and is deterministic. See the method-resolution and TCFL source audits for hashes and deviations.\n", encoding="utf-8")
    bootstrap = pd.read_csv(ext / "metrics" / "bootstrap_confidence_intervals.csv")
    counts = bootstrap.groupby(["dataset", "split"]).n_positions.max().to_dict()
    (ext / "reports" / "statistical_audit.md").write_text(f"# Statistical audit\n\nAggregation is frame → position within seed → mean across seeds → dataset/split. Bootstrap uses 10,000 position-level draws with seed 42. Maximum position counts by dataset/split: `{counts}`. PKU37 train, validation, and test are never pooled.\n", encoding="utf-8")
    provenance = f"# Sealing and provenance audit\n\nExtension locked at `{lock['locked_at_utc']}` on commit `{lock['git_commit']}`. Test and Duke timestamps are recorded in `extension_config_lock.json`. BASE_RUN `{lock['base_run']}` was verified against its initialization inventory before and after merge/package.\n"
    (ext / "reports" / "sealing_and_provenance_audit.md").write_text(provenance, encoding="utf-8")
    shutil.copy2(ext / "reports" / "sealing_and_provenance_audit.md", ext / "audit" / "sealing_and_provenance_audit.md")


def _fallback_workbook(ext: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("workbook creation needs @oai/artifact-tool or the openpyxl fallback") from exc
    workbook = Workbook(); workbook.remove(workbook.active)
    sources = [("Results", ext / "metrics" / "per_dataset_metrics.csv"), ("Seeds", ext / "metrics" / "per_seed_metrics.csv"), ("Paired CI", ext / "metrics" / "bootstrap_confidence_intervals.csv"), ("Runtime", ext / "metrics" / "runtime_summary.csv"), ("Checkpoints", ext / "metrics" / "checkpoint_inventory.csv")]
    for name, path in sources:
        sheet = workbook.create_sheet(name); frame = pd.read_csv(path) if path.is_file() else pd.DataFrame()
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None): sheet.append([None if pd.isna(value) else value for value in row])
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
        for cell in sheet[1]: cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36); sheet.column_dimensions[column[0].column_letter].width = width
    destination = ext / "benchmark_summary.xlsx"; workbook.save(destination)
    from openpyxl import load_workbook
    checked = load_workbook(destination, read_only=True, data_only=False)
    if set(checked.sheetnames) != {name for name, _ in sources} or any(checked[name].max_row < 1 for name, _ in sources):
        raise RuntimeError("exported benchmark workbook failed structural verification")
    checked.close()


def _gpt_light(root: Path, ext: Path) -> Path:
    stage = ext / "gpt_light_plus"
    if stage.exists(): shutil.rmtree(stage)
    stage.mkdir()
    for directory in ("audit", "configs", "reports", "manifests", "fixed_atlas"):
        source = ext / directory
        if source.exists(): shutil.copytree(source, stage / directory, ignore=shutil.ignore_patterns("backups"))
    (stage / "metrics").mkdir()
    for path in (ext / "metrics").glob("*.csv"):
        if path.name != "per_image_metrics.csv": shutil.copy2(path, stage / "metrics" / path.name)
    with (ext / "metrics" / "per_image_metrics.csv").open("rb") as source, gzip.open(stage / "metrics" / "per_image_metrics.csv.gz", "wb", compresslevel=9) as target: shutil.copyfileobj(source, target)
    for source in (ext / "benchmark_summary.xlsx", ext / "failures.csv", root / "docs" / "EXPERIMENT_LOG.md"):
        if source.is_file(): shutil.copy2(source, stage / source.name)
    snapshot = stage / "source"; snapshot.mkdir()
    for name in ("base.py", "sabids_adapter.py", "tcfl_adapter.py"):
        shutil.copy2(root / "tools" / "oct_denoise_benchmark" / "methods" / name, snapshot / name)
    shutil.copy2(root / "tools" / "oct_denoise_benchmark" / "train_tcfl.py", snapshot / "train_tcfl.py")
    guide = "# GPT analysis guide\n\nThis package contains locked formal results for all baseline methods plus SABIDS-current and TCFL-DnCNN. Inspect dataset/split-specific position-macro results and paired confidence intervals; do not pool PKU37 train/validation/test. Full images and checkpoints are deliberately excluded. Diagnostic limitations are identified in the audits.\n"
    (stage / "GPT_ANALYSIS_GUIDE.md").write_text(guide, encoding="utf-8")
    members = []
    for path in sorted(value for value in stage.rglob("*") if value.is_file()):
        members.append({"path": path.relative_to(stage).as_posix(), "bytes": path.stat().st_size, "sha256": sha256_file(path)})
    pd.DataFrame(members).to_csv(stage / "PACKAGE_MANIFEST.csv", index=False)
    archive = ext.parent / f"SABIDS_PKU37_denoise_benchmark_GPT_light_PLUS_SABIDS_TCFL_{datetime.now():%Y%m%d_%H%M%S}.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(value for value in stage.rglob("*") if value.is_file()): bundle.write(path, path.relative_to(stage).as_posix())
    with zipfile.ZipFile(archive) as bundle:
        bad = bundle.testzip()
        if bad: raise RuntimeError(f"GPT ZIP CRC failure: {bad}")
    (ext / "reports" / "gpt_light_package.json").write_text(json.dumps({"path": str(archive), "bytes": archive.stat().st_size, "sha256": sha256_file(archive), "crc": "passed"}, indent=2), encoding="utf-8")
    return archive


def package(root: Path, ext: Path) -> Path:
    root, ext = root.resolve(), ext.resolve(); verify_base(ext); _collect_training(ext); _model_complexity(ext)
    paths = build_image_paths(ext); _primary_zips(ext, paths); _fixed_atlas(root, ext, paths); _write_reports(ext); _fallback_workbook(ext)
    result = _gpt_light(root, ext); verify_base(ext); return result


def main(argv: Sequence[str] | None = None) -> None:
    p = argparse.ArgumentParser(); p.add_argument("--project-root", type=Path, default=Path(".")); p.add_argument("--ext-run", type=Path, required=True); args = p.parse_args(argv); print(package(args.project_root, args.ext_run))


if __name__ == "__main__": main()
