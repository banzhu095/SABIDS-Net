from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd


ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def write_workbook(path: str | Path, sheets: Mapping[str, pd.DataFrame], readme: list[list[object]] | None = None) -> Path:
    """Portable runtime writer used when Codex's artifact runtime is unavailable."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Excel export requires openpyxl in the ordinary Python runtime") from exc
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    if readme is not None:
        sheet = book.create_sheet("README")
        for row in readme: sheet.append(list(row))
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 100
        for row in sheet.iter_rows():
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet.max_row: sheet["A1"].font = Font(name="Arial", size=14, bold=True)
    for requested_name, frame in sheets.items():
        name = requested_name[:31]
        if name in book.sheetnames: raise ValueError(f"duplicate workbook sheet name: {name}")
        sheet = book.create_sheet(name)
        clean = frame.copy()
        clean = clean.where(pd.notna(clean), None)
        sheet.append(list(clean.columns))
        for row in clean.itertuples(index=False, name=None):
            sheet.append([value.item() if hasattr(value, "item") else value for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            maximum = max((len(str(cell.value or "")) for cell in column), default=0)
            sheet.column_dimensions[letter].width = min(max(maximum + 2, 10), 48)
            for cell in column[1:]:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=(maximum > 28 or "path" in str(column[0].value).lower()))
    temporary = destination.with_name(f".{destination.name}.tmp.xlsx")
    book.save(temporary)
    temporary.replace(destination)
    checked = load_workbook(destination, read_only=True, data_only=False)
    expected = (["README"] if readme is not None else []) + [name[:31] for name in sheets]
    if checked.sheetnames != expected: raise RuntimeError("workbook sheet verification failed")
    errors = []
    for sheet in checked.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = str(cell.value or "")
                if any(token in value for token in ERROR_TOKENS): errors.append(f"{sheet.title}!{cell.coordinate}:{value}")
    checked.close()
    if errors: raise RuntimeError("workbook formula error scan failed: " + "; ".join(errors[:20]))
    return destination


def highlight_group_best(path: str | Path, sheet_names: list[str], group_column: str = "tissue") -> None:
    """Highlight direction-aware best values only within a tissue/metric column."""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    high = {"psnr", "ssim", "epi", "local_contrast", "polarity_preserved"}
    low_tokens = ("rmse", "mae", "error", "abs_log")
    book = load_workbook(path)
    fill = PatternFill("solid", fgColor="C6EFCE")
    for name in sheet_names:
        if name not in book.sheetnames: continue
        sheet = book[name]
        headers = {str(cell.value): index for index, cell in enumerate(sheet[1], start=1)}
        if group_column not in headers: continue
        groups: dict[str, list[int]] = {}
        for row in range(2, sheet.max_row + 1): groups.setdefault(str(sheet.cell(row, headers[group_column]).value), []).append(row)
        for metric, column in headers.items():
            direction = "high" if metric in high else ("low" if any(token in metric for token in low_tokens) else None)
            if direction is None: continue
            for rows in groups.values():
                values = [(row, sheet.cell(row, column).value) for row in rows if isinstance(sheet.cell(row, column).value, (int, float))]
                if not values: continue
                best = (max if direction == "high" else min)(value for _, value in values)
                for row, value in values:
                    if value == best: sheet.cell(row, column).fill = fill
    temporary = Path(path).with_name(f".{Path(path).name}.highlight.tmp.xlsx")
    book.save(temporary); temporary.replace(path)
