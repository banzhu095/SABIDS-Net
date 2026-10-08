from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

from .common import read_json, require, sha, write_csv, write_json
from .data import Inputs


def atlas(run):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    run=Path(run)
    require((run/'per_frame_metrics.csv').exists(),'No completed test evaluation for atlas')
    config=read_json(run/'plan_lock.json')['config']
    data=pd.read_csv(run/'cohort.csv',keep_default_na=False)
    fixed=pd.read_csv(run/'fixed_atlas_samples.csv')
    folder=run/'atlas';folder.mkdir(exist_ok=True)
    for seed in config['segmentation_seeds']:
        for item in fixed.to_dict('records'):
            fig,axes=plt.subplots(len(config['methods']),10,figsize=(25,3*len(config['methods'])),squeeze=False)
            titles=['Noisy','Clean reference','Method input','GT layer','Predicted layer','GT vessel','Vessel probability','Binary vessel','TP green / FP red / FN blue','Choroid crop']
            reference=data[(data.sample_id==item['sample_id'])&(data.method_id=='clean_oracle')]
            # Clean reference is retained in audit even when the training arm is dropped.
            if reference.empty:
                audit=pd.read_csv(run/'input_asset_audit.csv',keep_default_na=False)
                reference=audit[(audit.sample_id==item['sample_id'])&(audit.method_id=='clean_oracle')]
            require(len(reference)==1 and bool(reference.iloc[0].available),'Missing clean atlas reference')
            clean=Inputs(reference.to_dict('records'),config['evaluation_size'],allow_test=True)[0][0][0].numpy()
            noisy_row=data[(data.sample_id==item['sample_id'])&(data.method_id=='noisy_identity')]
            noisy=Inputs(noisy_row.to_dict('records'),config['evaluation_size'],allow_test=True)[0][0][0].numpy()
            for i,method in enumerate(config['methods']):
                row=data[(data.method_id==method)&(data.sample_id==item['sample_id'])].to_dict('records')
                x,l,v,valid,vv,*_=Inputs(row,config['evaluation_size'],allow_test=True)[0]
                x,l,v=x[0].numpy(),l[0].numpy()>.5,v[0].numpy()>.5
                pred=np.load(run/'evaluation'/method/str(seed)/'predictions'/(item['sample_id']+'.npz'))
                lp,vp=pred['layer'],pred['vessel']; binary=vp>=.5
                color=np.zeros((*v.shape,3)); known=vv[0].numpy()>.5
                color[...,0]=binary&~v&known;color[...,1]=binary&v&known;color[...,2]=~binary&v&known
                scale=x.shape[0]/float(row[0]['height']); y=int(item['crop_y']*scale);xx=int(item['crop_x']*scale);s=max(1,int(item['crop_size']*scale))
                panels=[noisy,clean,x,l,lp>=.5,v,vp,binary,color,x[y:y+s,xx:xx+s]]
                for j,panel in enumerate(panels):
                    axes[i,j].imshow(panel,cmap='gray',vmin=0,vmax=1,interpolation='nearest')
                    axes[i,j].set_xticks([]);axes[i,j].set_yticks([])
                    if i==0:axes[i,j].set_title(titles[j],fontsize=9)
                axes[i,0].set_ylabel('SABIDS-current Stage-1/D0 output' if method=='sabids_current' else method,fontsize=9)
            fig.suptitle(f"{item['sample_id']}; seed {seed}; fixed threshold 0.5; common display [0,1]")
            fig.tight_layout();fig.savefig(folder/f"{item['sample_id']}_{item.get('role','middle')}_seed{seed}.png",dpi=120);plt.close(fig)


def build_report(run):
    run=Path(run);audit=read_json(run/'audit.json')
    config=read_json(run/'plan_lock.json')['config'] if (run/'plan_lock.json').exists() else {'mode':'formal'}
    if config['mode']=='formal':return formal_report(run,audit,config)
    complete=(run/'per_method_summary.csv').exists()
    status='completed' if complete else 'blocked_no_comparable_pilot_results'
    balance=pd.read_csv(run/'input_balance.csv')
    report='# Downstream segmentation pilot\n\nSABIDS-current Stage-1/D0 output; denoised-only information recoverability.\n\n'
    report+=f'Status: {status}. No Joint mechanism is evaluated.\n\n'+balance.to_markdown(index=False)+'\n\n'
    if complete:
        summary=pd.read_csv(run/'per_method_summary.csv');gains=pd.read_csv(run/'paired_position_gains.csv')
        report+=summary.to_markdown(index=False)+'\n\n'+gains.to_markdown(index=False)+'\n\n'
        report+='Preset conclusion: '+pd.read_csv(run/'pilot_success_criteria.csv').iloc[0]['status']+'\n\n'
        report+='SABIDS vs NAFNet vessel Dice, ROI Dice, Precision, Recall, outside-layer FP and boundary Dice are reported above by position. '
        if 'clean_oracle' in set(summary.method_id):
            table=summary.set_index('method_id')
            report+=f"Clean oracle minus Noisy vessel ROI Dice: {table.loc['clean_oracle','vessel_roi_dice_mean']-table.loc['noisy_identity','vessel_roi_dice_mean']:.6f}. "
        report+='Higher denoising PSNR alone does not determine downstream performance; this table tests that association in the fixed cohort.\n'
    else:
        report+='All eight scientific questions remain unanswered: no complete paired pilot has been trained or evaluated. '
        report+='Missing inputs/label disagreements are listed in missing_or_ambiguous_assets.csv. No advantage or disadvantage is inferred.\n'
        if not (run/'pilot_success_criteria.csv').exists():
            write_csv(run/'pilot_success_criteria.csv',[dict(status='not_evaluated',reason='Missing comparable pilot')])
    report+='\nFrames are aggregated within position, then within segmentation seed, then method. No significance p-values are reported for three positions. Boundary errors use pixels on the evaluation grid.\n'
    (run/'pilot_report.md').write_text(report,encoding='utf-8')
    count=next((x['positions'] for x in audit['labelled_counts'] if x['split']=='test'),0)
    formal='blocked_insufficient_labelled_test_positions' if count<6 else 'pending_full_asset_audit'
    (run/'formal_protocol.md').write_text(
        '# Formal protocol\n\nStatus: '+formal+'\n\nTen input methods; denoiser primary seed 42; segmentation seeds 42,123,2026. '
        'Shared existing NAF encoder and layer/vessel decoders without any denoising or interaction module. '
        'Training crops 512; evaluation 640; no method-specific optimization. '
        'All input assets and valid labels must pass audit. Complete pku_0025, pku_0038, pku_0043 labels if still missing. '
        'Three-position pilot is exploratory and is not formal paper evidence. '
        'Loss: masked layer BCE+Dice, vessel ROI BCE+Dice, outside-layer negative BCE weight 0.5. '
        'The legacy boundary head/containment auxiliary objective is omitted uniformly; this is an independent fixed segmenter, initialized once per seed. '
        'Pilot decision precedence: not_supportive, supportive, mixed-sign inconclusive, weak_support, otherwise inconclusive. '
        'Clearly worse FP means an absolute increase above 0.01; all thresholds fixed before test.\n',encoding='utf-8')
    runtimes=[]
    for p in (run/'tracks').glob('*/*/complete.json'):
        d=read_json(p);runtimes.append({k:d[k] for k in ['method_id','segmentation_seed','updates','seconds']})
    write_csv(run/'runtime_summary.csv',runtimes,['method_id','segmentation_seed','updates','seconds'])
    if not (run/'failures.csv').exists():write_csv(run/'failures.csv',[],['stage','reason'])
    tables={'Status':[{'item':'pilot','value':status},{'item':'formal','value':formal},{'item':'source manifest','value':audit['source_manifest']},
                      {'item':'meaning','value':'SABIDS-current Stage-1/D0 output'}, {'item':'evidence','value':'No performance claim before paired pilot completion'}],
            'Input balance':balance.to_dict('records')}
    for name,file in [('Position results','per_position_metrics.csv'),('Method results','per_method_summary.csv'),('Paired gains','paired_position_gains.csv'),('Runtime','runtime_summary.csv')]:
        if (run/file).exists():
            frame=pd.read_csv(run/file)
            if not frame.empty:tables[name]=json.loads(frame.to_json(orient='records'))
    write_json(run/'workbook_tables.json',tables)
    return status


def workbook(run):
    run=Path(run).resolve()
    node=os.environ.get('ARTIFACT_NODE') or shutil.which('node')
    builder=run/'build_summary.mjs'
    shutil.copyfile(Path(__file__).with_name('build_summary.mjs'),builder)
    formal=(run/'workbook_mode.json').exists() and read_json(run/'workbook_mode.json').get('formal')
    filename='benchmark_summary.xlsx' if formal else 'pilot_summary.xlsx'
    if node:
        before=(run/filename).stat().st_mtime_ns if (run/filename).exists() else 0
        result=subprocess.run([node,str(builder),str(run)],capture_output=True,text=True)
        write_json(run/'workbook_builder_log.json',dict(exit_code=result.returncode,stdout=result.stdout[-3000:],stderr=result.stderr[-3000:]))
        if result.returncode==0:
            return
        if (run/filename).is_file() and (run/filename).stat().st_mtime_ns>before and (run/'workbook_export_verified.json').exists():
            with zipfile.ZipFile(run/filename) as archive:
                require(archive.testzip() is None and 'xl/workbook.xml' in archive.namelist(),'Invalid exported workbook')
            write_json(run/'workbook_backend.json',dict(backend='artifact-tool',process_exit=result.returncode,verified_fresh_export=True))
            return
        require('ERR_MODULE_NOT_FOUND' in result.stderr, 'Workbook builder failed: '+result.stderr[-1500:])
    # Cloud without artifact-tool: standard portable XLSX fallback, explicitly logged.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook();wb.remove(wb.active)
    for name,rows in read_json(run/'workbook_tables.json').items():
        ws=wb.create_sheet(name);headers=list(rows[0]);ws.append(headers)
        for row in rows:ws.append([row.get(c) for c in headers])
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='34495E')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=28
        ws.freeze_panes='A2'
    wb.save(run/filename)
    write_json(run/'workbook_backend.json',dict(backend='openpyxl',reason='Node/artifact-tool unavailable'))


def package(run):
    run=Path(run);require((run/'pilot_summary.xlsx').exists() or (run/'benchmark_summary.xlsx').exists(),'Build workbook before packaging')
    snapshot=run/'source_snapshot';snapshot.mkdir(exist_ok=True)
    for path in Path(__file__).parent.glob('*.py'):shutil.copyfile(path,snapshot/path.name)
    files=[]
    candidates=[]
    for base,dirs,names in os.walk(run,followlinks=False):
        dirs[:]=[d for d in dirs if d not in {'predictions','node_modules','workbook_preview','materialized_images','exports','shared_initializations'}]
        candidates.extend(Path(base)/name for name in names)
    for path in candidates:
        if not path.is_file() or path.suffix.lower() not in {'.csv','.md','.json','.yaml','.xlsx','.png','.log','.py','.mjs'}:continue
        rel=path.relative_to(run)
        if any(part in {'predictions','node_modules','workbook_preview'} for part in rel.parts):continue
        if path.name in ['PACKAGE_MANIFEST.csv','package_verification.json']:continue
        files.append(path)
    manifest=[dict(path=p.relative_to(run).as_posix(),bytes=p.stat().st_size,sha256=sha(p)) for p in files]
    write_csv(run/'PACKAGE_MANIFEST.csv',manifest)
    archive=run/('GPT_light_'+run.name+'.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for path in files+[run/'PACKAGE_MANIFEST.csv']:z.write(path,path.relative_to(run).as_posix())
    with zipfile.ZipFile(archive) as z:
        require(z.testzip() is None,'ZIP CRC error')
        for row in manifest:require(hashlib.sha256(z.read(row['path'])).hexdigest()==row['sha256'],'ZIP SHA mismatch')
    result=dict(path=str(archive),sha256=sha(archive),files=len(manifest)+1,crc_passed=True)
    write_json(run/'package_verification.json',result)
    return result


def formal_report(run,audit,config):
    from .runner import completion_matrix
    matrix=pd.DataFrame(completion_matrix(run))
    balance=pd.read_csv(run/'input_balance.csv')
    complete=(run/'evaluation_status.json').exists()
    status=read_json(run/'evaluation_status.json')['status'] if complete else 'pending_formal_training_or_missing_assets'
    text='# Unified downstream segmentation\n\nSABIDS-current Stage-1/D0 denoised-only downstream segmentation. '
    text+='No Joint, D→S, S→D or UGBI mechanism is evaluated. TCFL inputs are unpaired denoising outputs.\n\n'
    text+=f'Status: {status}. Completed method × seed: {int((matrix.status=="complete").sum())}/30.\n\n'
    text+=balance.to_markdown(index=False)+'\n\n'
    text+='Canonical layer=(class 1|2), vessel=class 2, ignore=255. Fixed threshold 0.5. '
    text+='All methods share initialization, native patch plans, optimizer, budget, and full validation every five epochs. '
    text+='Checkpoint selection uses validation ROI Dice; differences ≤0.0001 tie-break on layer Dice, then earliest epoch.\n\n'
    text+='Frame → position within segmentation seed → seed mean per position → method. '
    text+='The 3 labelled test positions are the only independent units. CIs are descriptive clustered-position intervals, not significance tests. '
    text+='Annotate pku_0025, pku_0038, pku_0043 for evaluate-only expansion without retraining or checkpoint selection.\n\n'
    if complete:
        summary=pd.read_csv(run/'per_method_metrics.csv')
        columns=['method_id']+[m+'_mean' for m in ['layer_dice','vessel_dice','vessel_roi_dice','precision','recall','boundary_dice','vessel_outside_gt_layer_fp','thickness_mae']]
        text+=summary[columns].to_markdown(index=False)+'\n\n'
        gains=pd.read_csv(run/'paired_gains_by_position.csv')
        text+='## Position-paired effects\n\n'+gains[['comparison','position_id','vessel_roi_dice','recall','precision','boundary_dice','vessel_outside_gt_layer_fp']].to_markdown(index=False)+'\n\n'
        text+='Higher Recall accompanied by lower Precision, more outside-layer FP or lower Boundary Dice is consistent with oversegmentation, not an anatomical advantage. '
        text+='High-frequency energy near one must be interpreted alongside boundaries, precision and the fixed atlas. '
        text+='Method-level PSNR/SSIM–segmentation associations are descriptive; methods are not independent randomized observations.\n'
        denoise_join(run)
        plots(run)
    else:text+='Layer/Vessel performance, paired effects and PSNR–Dice relationships remain unavailable until all 30 checkpoints are locked and evaluated.\n'
    (run/'downstream_segmentation_report.md').write_text(text,encoding='utf-8')
    (run/'formal_protocol.md').write_text(text.split('## Position-paired effects')[0],encoding='utf-8')
    runtime=[]
    for p in (run/'tracks').glob('*/*/complete.json'):
        d=read_json(p);runtime.append({k:d[k] for k in ['method_id','segmentation_seed','updates','seconds']})
    write_csv(run/'runtime_summary.csv',runtime,['method_id','segmentation_seed','updates','seconds'])
    if not (run/'failures.csv').exists():write_csv(run/'failures.csv',[],['stage','reason'])
    tables={'Status':[dict(item='Experiment',value='Stage-1/D0 denoised-only downstream segmentation'),dict(item='Status',value=status),
        dict(item='Completed tracks',value=int((matrix.status=='complete').sum())),dict(item='Planned tracks',value=30),
        dict(item='Test inference',value='3 positions; descriptive CI only'),dict(item='Test threshold',value=.5)],
        'Input balance':balance.to_dict('records'),'Completion':matrix.to_dict('records')}
    for title,file in [('Methods','per_method_metrics.csv'),('Position effects','paired_gains_by_position.csv'),('Seed stability','per_seed_metrics.csv'),('Intervals','bootstrap_confidence_intervals.csv')]:
        if (run/file).exists():tables[title]=json.loads(pd.read_csv(run/file).to_json(orient='records'))
    write_json(run/'workbook_tables.json',tables);write_json(run/'workbook_mode.json',dict(formal=True))
    return status


def denoise_join(run):
    evidence=pd.read_csv(run/'formal_source_inventory.csv')
    tables=[]
    for row in evidence[evidence.kind=='per_image_metrics.csv'].to_dict('records'):
        if Path(row['path']).exists():tables.append(pd.read_csv(row['path'],keep_default_na=False))
    if not tables:
        write_csv(run/'denoise_segmentation_joined.csv',[],['method_id','position_id','reason'])
        write_json(run/'denoise_join_status.json',dict(status='missing_source_pixel_metrics'));return
    d=pd.concat(tables,ignore_index=True).drop_duplicates()
    if 'hf_energy_ratio_to_reference' in d:d['hf_energy_ratio']=d.hf_energy_ratio_to_reference
    d=d[(d.dataset=='PKU37')&(d.split=='test')]
    from .common import DEEP
    if 'seed' in d:d=d[(~d.method_id.isin(DEEP))|(pd.to_numeric(d.seed)==42)]
    metrics=[c for c in ['psnr','ssim','epi','reference_edge_mae','gradient_magnitude_mae','hf_energy_ratio','laplacian_energy_ratio'] if c in d]
    require(not d.duplicated(['method_id','sample_id']).any(),'Conflicting denoising metric sources')
    for c in metrics:d[c]=pd.to_numeric(d[c],errors='coerce')
    pos=pd.read_csv(run/'per_position_metrics.csv')
    d=d[d.position_id.isin(pos.position_id.unique())].groupby(['method_id','position_id'],as_index=False)[metrics].mean()
    joined=pos.merge(d,on=['method_id','position_id'],how='left',validate='one_to_one')
    write_csv(run/'denoise_segmentation_joined.csv',joined)
    method=joined.groupby('method_id').mean(numeric_only=True)
    from scipy.stats import spearmanr
    result=[]
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    for x,y in [('psnr','vessel_dice'),('ssim','vessel_dice'),('epi','boundary_dice'),('hf_energy_ratio','recall')]:
        if x not in method:continue
        points=method[[x,y]].replace([np.inf,-np.inf],np.nan).dropna()
        rho=float(spearmanr(points[x],points[y]).statistic) if len(points)>2 and points[x].nunique()>1 and points[y].nunique()>1 else None
        result.append(dict(x=x,y=y,methods=len(points),spearman=rho,scope='descriptive_no_p_value'))
        fig,ax=plt.subplots(figsize=(8,6));ax.scatter(points[x],points[y])
        for name,r in points.iterrows():ax.annotate(name,(r[x],r[y]),fontsize=8)
        ax.set_xlabel(x);ax.set_ylabel(y);fig.tight_layout();fig.savefig(run/(x+'_vs_'+y+'.png'),dpi=120);plt.close(fig)
    write_csv(run/'denoise_segmentation_spearman.csv',result)


def plots(run):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    folder=run/'training_curves';folder.mkdir(exist_ok=True)
    for p in (run/'tracks').glob('*/*/training_curve.csv'):
        d=pd.read_csv(p);fig,ax=plt.subplots(figsize=(7,4));ax.plot(d.epoch,d.loss,label='train loss')
        ax.plot(d.epoch,d.val_vessel_roi_dice,label='validation ROI Dice');ax.legend();ax.set_xlabel('epoch')
        fig.tight_layout();fig.savefig(folder/(p.parent.parent.name+'_seed'+p.parent.name+'.png'),dpi=120);plt.close(fig)
