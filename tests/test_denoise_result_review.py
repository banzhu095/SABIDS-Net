from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import yaml
import cv2
import numpy as np
import pytest
from openpyxl import load_workbook

from tools.denoise_result_review import METHOD_ORDER
from tools.denoise_result_review.image_io import decode_lossless, read_float01, to_float01
from tools.denoise_result_review.manifest_builder import primary_seeds
from tools.denoise_result_review.method_literature import implementation_summary, load_literature
from tools.denoise_result_review.run_discovery import inspect_run
from tools.denoise_result_review.workbook import write_workbook


def test_literature_covers_review_methods_and_separates_sources():
    frame = load_literature()
    assert set(frame.method_id) == set(METHOD_ORDER)
    tcfl = frame.set_index("method_id").loc["tcfl_dncnn"]
    assert tcfl.doi == "10.1109/TMI.2022.3184529"
    assert "unpaired" in tcfl.supervision
    assert pd.isna(frame.set_index("method_id").loc["sabids_current"].doi)


def test_discovery_rejects_structured_but_result_incomplete_run(tmp_path: Path):
    run = tmp_path / "runs" / "candidate"
    for directory in ("metrics", "configs", "audit", "manifests"): (run / directory).mkdir(parents=True, exist_ok=True)
    row = {"dataset": "PKU37", "split": "test", "position_id": "p1", "sample_id": "s1", "method_id": "noisy_identity", "seed": 0, "status": "success"}
    pd.DataFrame([row]).to_csv(run / "metrics" / "per_image_metrics.csv", index=False)
    for name in ("per_position_metrics.csv", "per_dataset_metrics.csv", "asset_inventory.csv", "selected_parameters.csv"): pd.DataFrame([row]).to_csv(run / "metrics" / name, index=False)
    pd.DataFrame([row]).to_csv(run / "manifests" / "denoised_dataset_manifest.csv", index=False)
    (run / "configs" / "inference_registry.yaml").write_text("methods: {}\n", encoding="utf-8")
    (run / "audit" / "config_lock.json").write_text(json.dumps({"status": "locked"}), encoding="utf-8")
    result = inspect_run(run)
    assert result.required_present == result.required_total
    assert result.package_ready is True and result.result_complete is False and result.pku_test_rows == 1
    assert "tcfl_dncnn" in result.missing_methods


def test_workbook_reopens_with_expected_sheets(tmp_path: Path):
    path = write_workbook(tmp_path / "中文结果.xlsx", {"Metrics": pd.DataFrame({"method_id": ["nlm"], "psnr": [30.5]})}, readme=[["Title", "Test"]])
    book = load_workbook(path, read_only=True, data_only=False)
    assert book.sheetnames == ["README", "Metrics"]
    assert book["Metrics"]["B2"].value == 30.5
    book.close()


def test_unicode_image_io_uses_fixed_uint_ranges_and_rejects_color(tmp_path: Path):
    folder = tmp_path / "中文 路径"; folder.mkdir()
    for dtype, maximum in ((np.uint8, 255), (np.uint16, 65535)):
        path = folder / f"image_{dtype.__name__}.png"
        image = np.array([[0, maximum]], dtype=dtype)
        ok, encoded = cv2.imencode(".png", image); assert ok; encoded.tofile(path)
        value, meta = read_float01(path)
        np.testing.assert_array_equal(value, np.array([[0.0, 1.0]], np.float32))
        assert meta["bit_depth"] == np.dtype(dtype).itemsize * 8
    fixed = to_float01(np.array([[100]], np.uint16))
    assert np.isclose(fixed.item(), 100 / 65535) and not np.isclose(fixed.item(), 1.0)
    with pytest.raises(ValueError, match="data_range"): to_float01(np.array([[0.5]], np.float32))
    color = folder / "color.png"; ok, encoded = cv2.imencode(".png", np.dstack([np.zeros((4,4), np.uint8), np.ones((4,4), np.uint8), np.zeros((4,4), np.uint8)])); assert ok; encoded.tofile(color)
    with pytest.raises(ValueError, match="non-grayscale"): decode_lossless(color)


def test_primary_deep_seed_comes_from_registry_not_metrics(tmp_path: Path):
    run = tmp_path / "run"; (run / "configs").mkdir(parents=True)
    (run / "configs" / "inference_registry.yaml").write_text("methods:\n  dncnn_paired: {seed: 123}\n  bm3d_standard: {seed: 0}\n", encoding="utf-8")
    seeds = primary_seeds(run)
    assert seeds["dncnn_paired"] == 123 and seeds["bm3d_standard"] == 0


def test_implementation_status_accepts_successful_per_image_evidence(tmp_path: Path):
    run = tmp_path / "run"; (run / "configs").mkdir(parents=True); (run / "metrics").mkdir()
    (run / "configs" / "inference_registry.yaml").write_text("methods: {}\n", encoding="utf-8")
    pd.DataFrame([{"method_id": "noisy_identity", "status": "success"}]).to_csv(run / "metrics" / "per_image_metrics.csv", index=False)
    summary = implementation_summary(tmp_path, run).set_index("method_id")
    assert summary.loc["noisy_identity", "status"] == "completed"
    assert summary.loc["noisy_identity", "completion_evidence"] == "successful_per_image"
    assert summary.loc["sabids_current", "status"] == "missing/not_completed"
